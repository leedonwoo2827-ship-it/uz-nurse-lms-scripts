# -*- coding: utf-8 -*-
"""모든 단계가 함께 쓰는 것 — 입력 시트 읽기, 팩 규격, 패널, 프롬프트 채우기, 모델 호출 기록.

규약
────────────────────────────────────────────────────────────────────────────
★ 이미 있는 산출 파일은 덮지 않는다(엔진 규약). 다시 만들려면 파일을 옮기거나 지운다.
★ 보낸 프롬프트는 전부 `_prompts/` 에 남긴다 — "왜 이렇게 나왔나" 에 답하기 위해.
★ 모델 응답은 JSON 또는 HTML 조각만 받는다. 앞뒤 잡담은 첫 `{`/`[`/`<` 앞을 버려 걷어낸다.
"""
from __future__ import annotations

import io
import json
import re
import sys
import time
from pathlib import Path

import yaml

ROOT = Path(__file__).resolve().parents[1]
PACK_DIR = ROOT / "pack"
PROMPT_DIR = ROOT / "tools" / "prompts"
DATA = ROOT / "data"
INPUT_XLSX = ROOT / "_context" / "차시설계_입력_v1.xlsx"
PANEL_JSON = DATA / "00_패널" / "패널.json"
LOG = DATA / "_log" / "run.jsonl"


def read(p: Path) -> str:
    return io.open(p, encoding="utf-8").read() if Path(p).is_file() else ""


def write(p: Path, text: str):
    Path(p).parent.mkdir(parents=True, exist_ok=True)
    io.open(p, "w", encoding="utf-8").write(text)


def site() -> dict:
    """이 과정이 어디를 위한 것인가 — 나라·기관·언어 이름. **저장소에 올리지 않는 파일.**

    ★ 이름을 한 곳에 모아 두는 이유는 두 가지다. 저장소를 공개해도 어느 사업인지
      드러나지 않고, 다른 나라·다른 분야로 옮길 때 이 파일만 바꾸면 된다.
    """
    p = PACK_DIR / "site.yaml"
    if not p.is_file():
        sys.exit("pack/site.yaml 이 없습니다.\n"
                 "  견본을 복사해서 채우십시오:  copy pack\\site.sample.yaml pack\\site.yaml")
    return yaml.safe_load(io.open(p, encoding="utf-8")) or {}


def pack() -> dict:
    """규격. 안에 {{country}} 같은 자리표시자가 있으면 site.yaml 값으로 채워 돌려준다."""
    raw = io.open(PACK_DIR / "pack.yaml", encoding="utf-8").read()
    for k, v in site().items():
        raw = raw.replace("{{%s}}" % k, str(v))
    return yaml.safe_load(raw) or {}


def terms_text() -> str:
    return "\n".join(l for l in read(PACK_DIR / "terms.csv").splitlines() if l.strip() and not l.startswith("#"))


# ── 입력 시트 ───────────────────────────────────────────────────────────────
def load_input():
    """`교과목`(코드→행) 과 `차시슬롯`(차시ID 순서 목록) — 사용자가 손본 판을 그대로 믿는다."""
    import openpyxl
    wb = openpyxl.load_workbook(INPUT_XLSX, data_only=True)
    subjects, slots = {}, []
    ws = wb["교과목"]
    cols = [c.value for c in ws[1]]
    for r in ws.iter_rows(min_row=2, values_only=True):
        if not r[cols.index("코드")]:
            continue
        row = {k: ("" if v is None else v) for k, v in zip(cols, r)}
        subjects[row["코드"]] = row
    ws = wb["차시슬롯"]
    cols = [c.value for c in ws[1]]
    for r in ws.iter_rows(min_row=2, values_only=True):
        if not r[cols.index("차시ID")]:
            continue
        slots.append({k: ("" if v is None else v) for k, v in zip(cols, r)})
    return subjects, slots


def slot(lecture_id: str):
    subjects, slots = load_input()
    for s in slots:
        if s["차시ID"] == lecture_id:
            return subjects[s["코드"]], s, [x for x in slots if x["코드"] == s["코드"]]
    sys.exit("차시슬롯에 없는 차시입니다: %s" % lecture_id)


# ── 패널 ────────────────────────────────────────────────────────────────────
def panel() -> dict:
    return json.load(io.open(PANEL_JSON, encoding="utf-8"))


def panel_cards(role: str) -> str:
    """프롬프트에 넣는 명단. role = author | reviewer."""
    out = []
    for p in panel()["panel"]:
        if p["role"] != role:
            continue
        head = "%s %s" % (p.get("번호", ""), p.get("가명", ""))
        if role == "author":
            out.append("- %s — %s(%s). 묻는 것: %s\n  약력: %s" % (
                head, p.get("view", ""), p["seat"], p.get("asks", ""), p.get("약력", "")))
        else:
            out.append("- %s — %s. %s\n  약력: %s" % (head, p["seat"], p.get("desc", ""), p.get("약력", "")))
    return "\n".join(out)


def panel_table_html() -> str:
    rows = []
    for p in panel()["panel"]:
        rows.append("<tr><td>%s</td><td>%s</td><td>%s</td><td>%s</td><td>%s</td></tr>" % (
            p.get("번호", ""), "저자진" if p["role"] == "author" else "검토진",
            p.get("view") or p["seat"], p.get("가명", ""), p.get("약력", "")))
    return ("<table><thead><tr><th>번호</th><th>구분</th><th>자리</th><th>가명</th><th>약력</th></tr></thead>"
            "<tbody>%s</tbody></table>" % "".join(rows))


# ── 프롬프트 ────────────────────────────────────────────────────────────────
def fill(template_name: str, **kw) -> str:
    tpl = read(PROMPT_DIR / template_name)
    if not tpl:
        sys.exit("프롬프트 파일이 없습니다: %s" % (PROMPT_DIR / template_name))
    cfg = pack()
    base = {
        "label": cfg.get("label", ""), "audience": cfg.get("audience", ""),
        "voice_tone": cfg["voice"]["tone"],
        "voice_forbid": ", ".join(cfg["voice"]["forbid"]),
        "voice_rules": "\n".join("- " + r for r in cfg["voice"]["rules"]),
        "localize_rules": "\n".join("- " + r for r in cfg["localize"]["rules"]),
        "localize_tag": cfg["localize"]["tag"],
        "screen_rules": "\n".join("- %s: %s" % (k, v) for k, v in cfg["screen"].items() if k != "svg_palette"),
        "svg_palette": ", ".join(cfg["screen"]["svg_palette"]),
        "terms": terms_text(),
        "say_min": cfg["say_total"][0], "say_max": cfg["say_total"][1],
        "slide_min": cfg["slides"][0], "slide_max": cfg["slides"][1],
        "per_min": cfg["say_per_slide"][0], "per_max": cfg["say_per_slide"][1],
        "block_min": cfg["body_block_chars"][0], "block_max": cfg["body_block_chars"][1],
        "sentence_max": cfg["sentence_max"], "new_terms_max": cfg["new_terms_max"],
        "quiz_min": cfg["quiz"][0], "quiz_max": cfg["quiz"][1],
        "duration": cfg["duration_min"],
    }
    base.update(kw)
    out = tpl
    for k, v in base.items():
        out = out.replace("{{%s}}" % k, str(v))
    # site.yaml 의 이름들(나라·기관·언어)은 마지막에 채운다 — 위에서 끼워 넣은 값 안에도 있을 수 있다.
    for k, v in site().items():
        out = out.replace("{{%s}}" % k, str(v))
    left = re.findall(r"\{\{(\w+)\}\}", out)
    if left:
        sys.exit("프롬프트에 안 채워진 자리: %s (%s)" % (", ".join(sorted(set(left))), template_name))
    return out


# ── 모델 호출 ───────────────────────────────────────────────────────────────
def call(prompt: str, log_to: Path, tag: str, retries: int = 6, wait: int = 600) -> str:
    """한 번 묻고 글을 받는다. 프롬프트·응답을 파일로 남긴다. 한도면 기다렸다 다시."""
    from _engine import ask, strip_fence, QuotaExceeded, NotAuthenticated, ProviderError
    log_to.mkdir(parents=True, exist_ok=True)
    write(log_to / (tag + "_프롬프트.txt"), prompt)
    t0 = time.time()
    for i in range(retries):
        try:
            text = strip_fence(ask(prompt, cwd=str(ROOT)))
            write(log_to / (tag + "_응답.txt"), text)
            _log({"tag": tag, "dir": str(log_to), "sec": round(time.time() - t0), "chars": len(text), "ok": True})
            return text
        except NotAuthenticated as e:
            _log({"tag": tag, "dir": str(log_to), "ok": False, "err": "auth"})
            sys.exit("[인증] %s\n→ VSCode 의 Claude Code 에서 로그인하십시오." % e)
        except QuotaExceeded as e:
            print("  [한도] %s — %d초 뒤 다시 (%d/%d)" % (str(e)[:80], wait, i + 1, retries))
            _log({"tag": tag, "dir": str(log_to), "ok": False, "err": "quota"})
            time.sleep(wait)
        except ProviderError as e:
            print("  [실패] %s — 60초 뒤 다시 (%d/%d)" % (str(e)[:120], i + 1, retries))
            _log({"tag": tag, "dir": str(log_to), "ok": False, "err": str(e)[:200]})
            time.sleep(60)
    raise RuntimeError("모델 호출이 %d번 실패했습니다: %s" % (retries, tag))


def _log(rec: dict):
    rec["t"] = time.strftime("%Y-%m-%d %H:%M:%S")
    LOG.parent.mkdir(parents=True, exist_ok=True)
    io.open(LOG, "a", encoding="utf-8").write(json.dumps(rec, ensure_ascii=False) + "\n")


def extract_json(text: str):
    """응답에서 JSON 하나를 꺼낸다. 앞뒤 잡담·펜스는 버린다."""
    t = text.strip()
    if t.startswith("```"):
        t = "\n".join(t.split("\n")[1:])
        if t.rstrip().endswith("```"):
            t = t.rstrip()[:-3]
    starts = [i for i in (t.find("{"), t.find("[")) if i >= 0]
    if not starts:
        raise ValueError("JSON 이 없습니다")
    s = min(starts)
    e = max(t.rfind("}"), t.rfind("]"))
    body = t[s:e + 1]
    try:
        return json.loads(body)
    except json.JSONDecodeError:
        pass
    # 흔한 결함: 뒤에 붙은 쉼표, 제어문자
    body2 = re.sub(r",\s*([}\]])", r"\1", body)
    body2 = re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", "", body2)
    try:
        return json.loads(body2, strict=False)
    except json.JSONDecodeError:
        pass
    # 문자열 안에 그대로 들어간 " (화면 HTML 의 "…" 인용) — 문맥으로 닫는 따옴표인지 판별해 복구
    return json.loads(repair_quotes(body2), strict=False)


def repair_quotes(s: str) -> str:
    """JSON 문자열 안의 이스케이프 안 된 " 를 \\" 로 바꾼다.
    닫는 따옴표 = 뒤에 공백을 건너뛰고 , } ] : 가 오는 것. 그 밖의 " 는 안쪽 따옴표로 본다."""
    out, i, n, in_str = [], 0, len(s), False
    while i < n:
        c = s[i]
        if not in_str:
            if c == '"':
                in_str = True
            out.append(c)
            i += 1
            continue
        if c == "\\" and i + 1 < n:
            out.append(c)
            out.append(s[i + 1])
            i += 2
            continue
        if c == '"':
            j = i + 1
            while j < n and s[j] in " \t\r\n":
                j += 1
            if j >= n or s[j] in ",}]:":
                in_str = False
                out.append(c)
            else:
                out.append('\\"')
            i += 1
            continue
        out.append(c)
        i += 1
    return "".join(out)


def say_len(s: str) -> int:
    return len((s or "").strip())
