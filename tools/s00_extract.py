# -*- coding: utf-8 -*-
"""0단계 — 원본 xlsx 에서 필요한 값만 뽑아 **병합 셀 없는** 입력 시트를 만든다.

    python tools/s00_extract.py                       → _context/차시설계_입력_v1.xlsx (+ .csv 2개)

무엇을 뽑는가
────────────────────────────────────────────────────────────────────────────
04·05 시트의 과목 행(코드 B-xx / R-xx 가 있는 행) + 06 시트(학습목표·교육내용·현지 유의)
+ 07 시트(콘텐츠 형태) 를 코드로 조인해 `교과목` 64행을 만들고, LMS credit 수만큼
`차시슬롯` 을 편다(최종평가는 0강). 이후 단계는 **이 파일만** 읽는다 — 원본은 다시 안 읽는다.

★ 이미 있는 출력 파일은 덮지 않는다. 사용자가 손본 판이 조용히 사라지면 안 된다.
  다시 만들려면 파일을 옮기거나 `--force`.
"""
from __future__ import annotations

import csv
import io
import sys
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from common import site

ROOT = Path(__file__).resolve().parents[1]
SRC = ROOT / "_context" / site()["source_xlsx"]
OUT = ROOT / "_context" / "차시설계_입력_v1.xlsx"

COURSE_SHEETS = [("기본", "04_기본 간호 실무"), ("응급", "05_응급상황에서의 간호")]
BATCH1_PER_COURSE = 30          # 1차 배치 = 과정마다 앞쪽 이만큼

SUBJ_COLS = ["과정", "모듈코드", "모듈명", "코드", "교과목명", "주역량", "부역량",
             "credit", "LMS", "집합", "운영", "교육방법", "평가방법", "벤치마킹 근거",
             "학습목표", "주요 교육내용", "현지 적용 유의", "콘텐츠 형태",
             "강의수", "1차배치 강의수", "비고"]
SLOT_COLS = ["차시ID", "과정", "모듈코드", "모듈명", "코드", "교과목명", "차수", "총차수",
             "배치", "제목", "학습목표", "상태", "비고"]


def s(v):
    return "" if v is None else str(v).strip()


def read_course(wb, course, sheet):
    ws = wb[sheet]
    rows, module_code, module_name = [], "", ""
    for r in ws.iter_rows(min_row=5, values_only=True):
        a = s(r[0])
        code = s(r[1])
        # 모듈 머리행: 'B-M1. 전문직 기초와 법적 책무 (16 credit) ◆ 차용 : …'
        if a and not code and a[:1] in "BR" and "-M" in a[:5]:
            head = a.split("◆")[0].strip()
            module_code = head.split(".")[0].strip()
            module_name = head.split(".", 1)[1].strip() if "." in head else ""
            # '(16 credit)' 꼬리 제거
            if "(" in module_name:
                module_name = module_name[:module_name.rfind("(")].strip()
            continue
        if not (code.startswith("B-") or code.startswith("R-")) or code == a:
            continue
        if not s(r[2]) or s(r[2]) == "합계":
            continue
        rows.append({
            "과정": course, "모듈코드": a or module_code, "모듈명": module_name,
            "코드": code, "교과목명": s(r[2]), "주역량": s(r[3]), "부역량": s(r[4]),
            "credit": r[5] or 0, "LMS": r[6] or 0, "집합": r[7] or 0,
            "운영": s(r[8]), "교육방법": s(r[9]), "평가방법": s(r[10]),
            "벤치마킹 근거": s(r[11]),
        })
    return rows


def read_detail(wb):
    ws = wb["06_교과목별 상세"]
    out = {}
    for r in ws.iter_rows(min_row=5, values_only=True):
        code = s(r[0])
        if code.startswith(("B-", "R-")) and s(r[1]):
            out[code] = {"학습목표": s(r[5]), "주요 교육내용": s(r[6]),
                         "현지 적용 유의": s(r[8])}
    return out


def read_content_type(wb):
    ws = wb["07_LMS·집합 운영 설계"]
    out = {}
    for r in ws.iter_rows(min_row=1, values_only=True):
        code = s(r[1])
        if code.startswith(("B-", "R-")) and s(r[4]) and "credit" not in s(r[3]):
            # 첫 표(LMS)만. 두 번째 표(집합)는 5열이 기자재라 '동영상/자율학습' 낱말이 없다.
            if "동영상" in s(r[4]) or "자율학습" in s(r[4]):
                out.setdefault(code, s(r[4]))
    return out


def build(force=False):
    if OUT.exists() and not force:
        sys.exit("이미 있습니다: %s\n덮지 않았습니다. 다시 만들려면 파일을 옮기거나 --force." % OUT)
    wb = openpyxl.load_workbook(SRC, data_only=True)
    detail, ctype = read_detail(wb), read_content_type(wb)

    subjects, slots = [], []
    for course, sheet in COURSE_SHEETS:
        rows = read_course(wb, course, sheet)
        acc = 0
        for row in rows:
            row.update(detail.get(row["코드"], {"학습목표": "", "주요 교육내용": "", "현지 적용 유의": ""}))
            row["콘텐츠 형태"] = ctype.get(row["코드"], "")
            is_final = row["교과목명"].strip() == "최종평가"
            n = 0 if is_final else int(row["LMS"] or 0)
            row["강의수"] = n
            b1 = max(0, min(n, BATCH1_PER_COURSE - acc))
            row["1차배치 강의수"] = b1
            row["비고"] = "최종평가 — LMS credit 은 자율학습 과제, 강의 없음" if is_final else ""
            acc += n
            subjects.append(row)
            for k in range(1, n + 1):
                slots.append({
                    "차시ID": "%s-%d" % (row["코드"], k), "과정": course,
                    "모듈코드": row["모듈코드"], "모듈명": row["모듈명"],
                    "코드": row["코드"], "교과목명": row["교과목명"],
                    "차수": k, "총차수": n, "배치": 1 if k <= b1 else 2,
                    "제목": "", "학습목표": "", "상태": "미설계", "비고": "",
                })

    out = openpyxl.Workbook()
    ws = out.active
    ws.title = "교과목"
    write_table(ws, SUBJ_COLS, subjects)
    ws2 = out.create_sheet("차시슬롯")
    write_table(ws2, SLOT_COLS, slots)
    ws3 = out.create_sheet("읽는법")
    notes = [
        ("이 파일", "원본 xlsx(04·05·06·07 시트)에서 필요한 값만 뽑아 조인한 것. 병합 셀 없음. 이후 모든 단계는 이 파일만 읽는다."),
        ("교과목 시트", "과목 1행. 학습목표·주요 교육내용·현지 적용 유의 = 06 시트, 콘텐츠 형태 = 07 시트."),
        ("강의수", "= LMS credit. 1 credit = 45분 동영상 1강. 최종평가 과목은 0."),
        ("1차배치 강의수", "과정마다 과목 순으로 누적 %d강까지가 1차. 나머지가 2차." % BATCH1_PER_COURSE),
        ("차시슬롯 시트", "강의 1행. 차시ID = 코드-차수. 제목·학습목표는 2단계(설계)가 채운다. 손으로 미리 적어도 된다 — 설계가 그것을 존중한다."),
        ("배치", "1 = 1차, 2 = 나머지."),
        ("상태", "미설계 → 설계 → 초고 → 검토 → 탈고. 도구가 갱신하지 않는다(카탈로그 index.html 이 실제 상태를 보여 준다)."),
        ("수정 규칙", "행을 지우거나 차시ID 를 바꾸면 그 강은 만들지 않는다. 강의수를 바꾸면 차시슬롯 행도 맞춰 늘리거나 줄인다."),
        ("합계", site()["course_summary"]),
    ]
    write_table(ws3, ["항목", "내용"], [dict(항목=a, 내용=b) for a, b in notes])
    ws3.column_dimensions["B"].width = 110
    out.save(OUT)

    for name, cols, rows in (("교과목", SUBJ_COLS, subjects), ("차시슬롯", SLOT_COLS, slots)):
        p = OUT.with_name("차시설계_입력_v1_%s.csv" % name)
        with io.open(p, "w", encoding="utf-8-sig", newline="") as f:
            w = csv.DictWriter(f, fieldnames=cols)
            w.writeheader()
            w.writerows(rows)

    by = {}
    for r in slots:
        by.setdefault(r["과정"], [0, 0])
        by[r["과정"]][r["배치"] - 1] += 1
    print("교과목 %d행 · 차시슬롯 %d행 → %s" % (len(subjects), len(slots), OUT))
    for c, (a, b) in by.items():
        print("  %s: 1차 %d + 2차 %d = %d강" % (c, a, b, a + b))
    return 0


def write_table(ws, cols, rows):
    ws.append(cols)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F4E79")
        c.alignment = Alignment(vertical="center", wrap_text=True)
    for r in rows:
        ws.append([r.get(k, "") for k in cols])
    for i, k in enumerate(cols, 1):
        width = 14
        if k in ("교과목명", "모듈명", "제목"): width = 34
        if k in ("벤치마킹 근거", "학습목표", "주요 교육내용", "현지 적용 유의"): width = 60
        ws.column_dimensions[get_column_letter(i)].width = width
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(vertical="top", wrap_text=True)
    ws.freeze_panes = "A2"


if __name__ == "__main__":
    sys.exit(build(force="--force" in sys.argv))
