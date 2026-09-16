# -*- coding: utf-8 -*-
"""1단계 — 저자진 4 + 검토진 16 패널을 뽑고 **엑셀 명부**로 남긴다.

    python tools/s01_panel.py            pack/persona.yaml → 엔진 sample_panel → data/00_패널/
    python tools/s01_panel.py --xlsx     이미 뽑은 json 으로 명부만 다시 만든다

엔진(260810-MatrAIx-검토분석엔진)의 tools/sample_panel.py 를 그대로 쓴다.
같은 씨앗·같은 조각이면 같은 16명이 다시 나온다(재현 가능성이 이 명부의 근거다).
"""
from __future__ import annotations

import io
import json
import shutil
import subprocess
import sys
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from _engine import ENGINE, VENV_PY
from common import site

ROOT = Path(__file__).resolve().parents[1]
PACK = "uz-nurse"
UNIT = "PANEL"
OUT_DIR = ROOT / "data" / "00_패널"
JSON = OUT_DIR / "패널.json"
XLSX = OUT_DIR / "패널명부.xlsx"
MD = OUT_DIR / "패널명부.md"

PROFILE = [("age_bracket", "나이"), ("region", "지역"), ("lang_korean", "한국어"),
           ("english_proficiency", "영어"), ("ind_healthcare", "보건의료 경력"),
           ("fam_nursing", "간호 친숙도"), ("fam_pharmacology", "약리 친숙도"),
           ("role_function", "직능"), ("domain", "분야"),
           ("domain_characteristics", "성향"), ("highest_education", "학력"),
           ("years_experience", "경력 연수"), ("skill_public_speaking", "발표 역량")]


def sample():
    eng_pack = ENGINE / "packs" / PACK
    eng_pack.mkdir(parents=True, exist_ok=True)
    shutil.copy(ROOT / "pack" / "persona.yaml", eng_pack / "persona.yaml")
    if not (eng_pack / "pack.yaml").is_file():
        io.open(eng_pack / "pack.yaml", "w", encoding="utf-8").write(
            "# 패널 추출 전용 팩. 집필 파이프라인은 %s 에 있다.\nslug: %s\nlabel: %s\npersona: persona.yaml\n"
            % (ROOT, PACK, site()["label"]))
    (ENGINE / "data" / PACK / UNIT).mkdir(parents=True, exist_ok=True)
    src = ENGINE / "data" / PACK / UNIT / "01_선정" / ("01_%s_패널.json" % UNIT)
    if src.exists():
        src.unlink()          # 샘플러는 덮지 않으므로 우리가 치운다(조건 파일이 바뀌었을 수 있다)
    r = subprocess.run([str(VENV_PY), "tools/sample_panel.py", PACK, UNIT], cwd=str(ENGINE),
                       env={**__import__("os").environ, "PYTHONIOENCODING": "utf-8"})
    if r.returncode != 0 or not src.exists():
        sys.exit("패널 추출 실패 (엔진 sample_panel.py)")
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    shutil.copy(src, JSON)
    print("→", JSON)


def load_seats():
    import yaml
    cfg = yaml.safe_load(io.open(ROOT / "pack" / "persona.yaml", encoding="utf-8"))
    seats = {s["id"]: s for s in cfg["author"]["seats"]}
    strata = {s["id"]: s for s in cfg["reviewer"]["strata"]}
    return cfg, seats, strata


def roster():
    d = json.load(io.open(JSON, encoding="utf-8"))
    cfg, seats, strata = load_seats()
    rows = []
    n_auth = n_rev = 0
    for p in d["panel"]:
        c = p["persona"]
        dims = c.get("dimensions", {})
        if p["role"] == "author":
            n_auth += 1
            no = "A%d" % n_auth
            s = seats.get(p["seat"], {})
            brief = "관점: %s / 묻는 것: %s" % (p.get("view", ""), p.get("asks", ""))
        else:
            n_rev += 1
            no = "R%02d" % n_rev
            s = strata.get(p["seat"], {})
            brief = s.get("desc", "")
        row = {"번호": no, "구분": "저자진" if p["role"] == "author" else "검토진",
               "자리": p["seat"], "가명": p.get("가명", ""), "약력": p.get("약력", ""),
               "역할 설명": brief,
               "페르소나 ID": c.get("id", ""), "출처": c.get("source", ""),
               "접지": c.get("grounding", ""), "통과 조건": c.get("gate", ""),
               "채워진 차원 수": c.get("populated_dimensions", "")}
        for k, label in PROFILE:
            row[label] = dims.get(k, "")
        row["그 밖의 차원"] = "; ".join("%s=%s" % (k, v) for k, v in sorted(dims.items())
                                   if k not in dict(PROFILE))
        row["참여 강의"] = ""
        row["비고"] = "가상 인물. 실제 인물이 아니다."
        rows.append(row)
    cols = list(rows[0].keys())

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "명부"
    ws.append(cols)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F4E79")
    for r in rows:
        ws.append([r[k] for k in cols])
    for i, k in enumerate(cols, 1):
        ws.column_dimensions[get_column_letter(i)].width = 60 if k in ("역할 설명", "그 밖의 차원", "약력", "참여 강의") else 16
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(vertical="top", wrap_text=True)
    ws.freeze_panes = "C2"

    ws2 = wb.create_sheet("선발 조건")
    ws2.append(["구분", "자리", "인원", "must", "must_any", "prefer", "설명"])
    for c in ws2[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F4E79")
    ws2.append(["공통", "-", "", "", json.dumps(cfg["common"].get("must_any"), ensure_ascii=False), "", "보건의료를 아는 사람"])
    ac = cfg["author"].get("common", {})
    ws2.append(["저자진 공통", "-", 4, json.dumps(ac.get("must"), ensure_ascii=False), "", json.dumps(ac.get("prefer"), ensure_ascii=False), cfg["author"].get("note", "")])
    for s in cfg["author"]["seats"]:
        ws2.append(["저자진", s["id"] + " " + s["view"], 1, json.dumps(s.get("must"), ensure_ascii=False),
                    json.dumps(s.get("must_any"), ensure_ascii=False), json.dumps(s.get("prefer"), ensure_ascii=False), s.get("asks", "")])
    for s in cfg["reviewer"]["strata"]:
        ws2.append(["검토진", s["id"], s["n"], json.dumps(s.get("must"), ensure_ascii=False),
                    json.dumps(s.get("must_any"), ensure_ascii=False), json.dumps(s.get("prefer"), ensure_ascii=False), s.get("desc", "")])
    ws2.append([])
    ws2.append(["재현", "python tools/s01_panel.py", "", "seed=%s" % d.get("seed"), "", "shards=%s" % ", ".join(d.get("shards", [])), d.get("dataset", "")])
    ws2.append(["한계", cfg.get("limits", "").strip()])
    for i in range(1, 8):
        ws2.column_dimensions[get_column_letter(i)].width = 40
    for row in ws2.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(vertical="top", wrap_text=True)
    wb.save(XLSX)

    md = ["# 패널 명부 — 저자진 %d · 검토진 %d (씨앗 %s)" % (n_auth, n_rev, d.get("seed")), "",
          "| 번호 | 구분 | 자리 | 가명 | 약력 | 나이 | 지역 | 보건의료 | 간호 | 직능 | 학력 |", "|---|---|---|---|---|---|---|---|---|---|---|"]
    for r in rows:
        md.append("| %s | %s | %s | %s | %s | %s | %s | %s | %s | %s | %s |" % (
            r["번호"], r["구분"], r["자리"], r["가명"], r["약력"], r["나이"], r["지역"], r["보건의료 경력"],
            r["간호 친숙도"], r["직능"], r["학력"]))
    md += ["", "가상의 응답자다. 실제 수강생·전문가 데이터가 아니다. " + cfg.get("limits", "").strip()]
    io.open(MD, "w", encoding="utf-8").write("\n".join(md) + "\n")
    print("→", XLSX)
    print("→", MD)
    print("저자진 %d · 검토진 %d" % (n_auth, n_rev))


if __name__ == "__main__" and "--bio" not in sys.argv:
    if "--xlsx" not in sys.argv:
        sample()
    roster()


# ── 약력 — 모델이 한 번 쓴다 ─────────────────────────────────────────────────
BIO_PROMPT = """아래는 {country} {learner} 대상 LMS 강의(45분 × {total}강)를 만들기 위해
페르소나 데이터셋에서 뽑은 저자진 4명과 검토진 16명의 차원값입니다. 실제 인물이 아닌 가상의 응답자입니다.

각 사람에게 **가명**(지역에 맞는 이름; 동아시아면 한국 이름, {regions}면 {country}에서 흔한 이름)과
**약력 2~3문장**(한국어, 존칭 없이 담담하게)을 써 주십시오. 약력은 차원값과 모순되지 않아야 하고,
이 과제에서 맡은 자리(저자진의 관점 / 검토진의 층)에 맞아야 합니다. 없는 수치나 기관명을 지어내지 마십시오 —
"타슈켄트의 한 지역병원" 처럼 일반적으로 적습니다. 검토진 학습자 층(1년차·2·3년차)은 그 연차에 맞는 경험을 적습니다.

응답은 JSON 하나만. 앞뒤 설명 없이. 형식:
{"A1": {"가명": "...", "약력": "..."}, "R01": {...}, ...}

## 명단
%s
"""


def bio():
    from _engine import ask, strip_fence
    d = json.load(io.open(JSON, encoding="utf-8"))
    cards, n_auth, n_rev = [], 0, 0
    for p in d["panel"]:
        if p["role"] == "author":
            n_auth += 1; no = "A%d" % n_auth
            head = "저자진 %s · 관점: %s · 묻는 것: %s" % (p["seat"], p.get("view", ""), p.get("asks", ""))
        else:
            n_rev += 1; no = "R%02d" % n_rev
            head = "검토진 %s · %s" % (p["seat"], p.get("desc", ""))
        dims = "; ".join("%s=%s" % kv for kv in sorted(p["persona"].get("dimensions", {}).items()))
        cards.append("- %s: %s\n  차원: %s" % (no, head, dims))
    s = site()
    prompt = (BIO_PROMPT.replace("{country}", s["country"]).replace("{learner}", s["learner"])
              .replace("{total}", str(s["total_lectures"])).replace("{regions}", s["name_regions"])
              % "\n".join(cards))
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    io.open(OUT_DIR / "약력_프롬프트.txt", "w", encoding="utf-8").write(prompt)
    print("모델을 부릅니다(약력 20명, 1회)...")
    text = strip_fence(ask(prompt, cwd=str(ROOT)))
    text = text[text.find("{"):text.rfind("}") + 1]
    bios = json.loads(text)
    n_auth = n_rev = 0
    for p in d["panel"]:
        if p["role"] == "author":
            n_auth += 1; no = "A%d" % n_auth
        else:
            n_rev += 1; no = "R%02d" % n_rev
        p["번호"] = no
        p["가명"] = bios.get(no, {}).get("가명", "")
        p["약력"] = bios.get(no, {}).get("약력", "")
    io.open(JSON, "w", encoding="utf-8").write(json.dumps(d, ensure_ascii=False, indent=1))
    print("약력 %d명 기록 →" % len(bios), JSON)


if __name__ == "__main__" and "--bio" in sys.argv:
    bio()
    roster()
