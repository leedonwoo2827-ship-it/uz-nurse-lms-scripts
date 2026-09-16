# -*- coding: utf-8 -*-
"""10단계 — 전체 현황(도구, 모델 없음). 언제든 다시 돌린다.

    python tools/s10_catalog.py
      → data/index.html          전 강 표: 제목·장수·글자수·분·상태·링크. 분량 OK 판단은 여기서.
      → data/검토의견.xlsx        강마다 시트 1개(122개까지): 이상무·지적(AS IS/TO BE·선택 드롭다운) + 맨 끝 패널 20명
      → data/00_패널/패널명부.xlsx 의 '참여 강의' 열 갱신
      → data/목차.xlsx            설계된 강 제목·학습목표 (차시슬롯 순)
"""
from __future__ import annotations

import html
import io
import json
import re
import sys
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from common import DATA, load_input, pack, read, site

LECT = DATA / "02_강의"
HEAD_FILL = PatternFill("solid", fgColor="1F4E79")
HEAD_FONT = Font(bold=True, color="FFFFFF")
OK_FILL = PatternFill("solid", fgColor="E8F5E9")
ISSUE_FILL = {"상": PatternFill("solid", fgColor="FFCDD2"), "중": PatternFill("solid", fgColor="FFF3CD"), "하": PatternFill("solid", fgColor="F1F5F9")}


def status_of(lid: str) -> dict:
    d = LECT / lid
    st = {"설계": False, "스토리보드": False, "초고": False, "검토": False, "탈고": False,
          "장수": "", "글자수": "", "분": "", "판정": "", "탈고판정": "", "지적": "", "제목": ""}
    code = lid.rsplit("-", 1)[0]
    dp = DATA / "01_설계" / ("%s_설계.json" % code)
    if dp.is_file():
        st["설계"] = True
        for l in json.load(io.open(dp, encoding="utf-8")).get("lectures", []):
            if l["id"] == lid:
                st["제목"] = l.get("title", "")
                st["학습목표"] = l.get("objectives", [])
    st["스토리보드"] = (d / "04_스토리보드.json").is_file()
    st["초고"] = (d / ("%s_원고.html" % lid)).is_file()
    ck = d / "07_검수.json"
    if ck.is_file():
        c = json.load(io.open(ck, encoding="utf-8"))
        st.update({"장수": c.get("slides"), "글자수": c.get("total_say"), "분": c.get("minutes"), "판정": c.get("status")})
    rv = d / "08_검토.json"
    if rv.is_file():
        st["검토"] = True
        r = json.load(io.open(rv, encoding="utf-8"))
        st["지적"] = sum(1 for it in r.get("items", []) if it.get("kind") != "이상무")
    st["탈고"] = (d / ("%s_탈고.html" % lid)).is_file()
    ck2 = d / ("%s_검수표_탈고.md" % lid)
    if ck2.is_file():
        m = re.search(r"판정: \*\*(.+?)\*\*", read(ck2))
        st["탈고판정"] = m.group(1) if m else ""
    return st


def index_html(slots, stats):
    cfg = pack()
    rows = []
    tot = {"초고": 0, "탈고": 0, "글자": 0, "n": 0}
    for s in slots:
        lid = s["차시ID"]; st = stats[lid]
        link = lambda name: ('<a href="02_강의/%s/%s">%s</a>' % (lid, "%s_%s.html" % (lid, name), name)) if (LECT / lid / ("%s_%s.html" % (lid, name))).is_file() else ""
        stage = "탈고" if st["탈고"] else "검토" if st["검토"] else "초고" if st["초고"] else "스토리보드" if st["스토리보드"] else "설계" if st["설계"] else "-"
        if st["초고"]: tot["초고"] += 1
        if st["탈고"]: tot["탈고"] += 1
        if st["글자수"]: tot["글자"] += int(st["글자수"]); tot["n"] += 1
        cls = "ok" if st["판정"] == "통과" else ("hold" if st["판정"] else "")
        rows.append("<tr class='%s'><td>%s</td><td>%s</td><td>%s</td><td>%s</td><td>%s</td><td>%s</td><td>%s</td><td>%s</td><td>%s</td><td>%s</td><td>%s</td><td>%s %s</td></tr>" % (
            cls, s["배치"], s["과정"], s["모듈코드"], lid, html.escape(s["교과목명"]), html.escape(st["제목"] or s.get("제목") or ""),
            stage, st["장수"] or "", ("{:,}".format(st["글자수"]) if st["글자수"] else ""), st["분"] or "",
            html.escape(str(st["판정"] or "")), link("원고"), link("탈고")))
    head = """<!DOCTYPE html><html lang="ko"><head><meta charset="utf-8"><title>@T — 현황</title>
<style>body{font-family:"Noto Sans KR","Malgun Gothic",sans-serif;font-size:13px;color:#334155;padding:16px 24px}
table{border-collapse:collapse;width:100%}th{background:#1F4E79;color:#fff;padding:6px 8px;text-align:left;position:sticky;top:0}
td{border-bottom:1px solid #E2E8F0;padding:5px 8px;vertical-align:top}tr.ok td{background:#F1F8F4}tr.hold td{background:#FFF8E1}
a{color:#2E75B6}.sum{margin:0 0 14px;color:#64748B}</style></head><body>
<h1>@T — 현황</h1>
<p class="sum">초고 @A / 탈고 @B / @N강 · 초고 평균 낭독 @C자 (규격 @D~@E자, 5.5자/초 → 45분 본문) · 규격: 슬라이드 @F~@G장</p>
<p class="sum">파일: <a href="검토의견.xlsx">검토의견.xlsx</a> (강별 시트) · <a href="목차.xlsx">목차.xlsx</a> · <a href="00_패널/패널명부.xlsx">패널명부.xlsx</a></p>
<table><thead><tr><th>배치</th><th>과정</th><th>모듈</th><th>차시ID</th><th>교과목</th><th>강 제목</th><th>단계</th><th>장</th><th>낭독자</th><th>분</th><th>검수</th><th>파일</th></tr></thead><tbody>
"""
    for k, v in (("@T", site()["label"]), ("@N", len(slots)), ("@A", tot["초고"]), ("@B", tot["탈고"]), ("@C", "{:,}".format(tot["글자"] // tot["n"]) if tot["n"] else "-"),
                 ("@D", "{:,}".format(cfg["say_total"][0])), ("@E", "{:,}".format(cfg["say_total"][1])),
                 ("@F", cfg["slides"][0]), ("@G", cfg["slides"][1])):
        head = head.replace(k, str(v))
    return head + "\n".join(rows) + "\n</tbody></table></body></html>\n"


def review_xlsx(slots, stats):
    cols = ["장#", "장ID", "코너", "장 제목", "구분", "검토자", "층", "항목", "AS IS", "TO BE", "근거", "심각도", "확신", "선택", "반영"]
    wb = openpyxl.Workbook()
    ws0 = wb.active
    ws0.title = "읽는법"
    for r in [("이 파일", "강마다 시트 하나. 검토진 16명의 의견을 **의견으로** 남긴 것. 결론(탈고 HTML)은 따로 완결되어 있다."),
              ("구분", "이상무 = 읽었고 문제 없음(무엇을 이해했는지 근거 열에 한 줄). 지적 = AS IS(원문 인용) / TO BE(고친 문장)."),
              ("선택", "드롭다운: 1안 적용 / 기각 / 보류. 도구가 '상·중 + 확신' 이면 1안 적용으로 미리 골라 두었다. 사람이 바꾼다."),
              ("반영", "탈고에 반영(1안) / 저자진 미반영 / 미반영 / 실패. 도구가 적는다. 사람이 선택을 바꾸면 s09 를 다시 돌린다(--force)."),
              ("맨 끝", "그 강의 집필·검토 패널 20명(가명·약력). 강별로 다를 수 있어 강마다 둔다."),
              ("색", "빨강 = 심각도 상, 노랑 = 중, 회색 = 하, 초록 = 이상무.")]:
        ws0.append(r)
    ws0.column_dimensions["A"].width = 12; ws0.column_dimensions["B"].width = 120
    made = 0
    for s in slots:
        lid = s["차시ID"]
        p = LECT / lid / "08_검토.json"
        if not p.is_file():
            continue
        r = json.load(io.open(p, encoding="utf-8"))
        ws = wb.create_sheet(lid[:31])
        ws.append(["%s %s" % (lid, r.get("title", ""))]); ws["A1"].font = Font(bold=True, size=13)
        ws.append(["요약", r.get("summary", "")])
        ws.append(["층별 이해도(0~5)", "; ".join("%s %s" % kv for kv in (r.get("scores") or {}).items())])
        ws.append(["지적 %d건 · 이상무 %d행" % (sum(1 for it in r["items"] if it.get("kind") != "이상무"), sum(1 for it in r["items"] if it.get("kind") == "이상무"))])
        ws.append([])
        ws.append(cols)
        hdr = ws.max_row
        for c in ws[hdr]:
            c.font = HEAD_FONT; c.fill = HEAD_FILL
        for it in r["items"]:
            ws.append([it.get("no"), it.get("id"), it.get("corner"), it.get("slide_title"),
                       "이상무" if it.get("kind") == "이상무" else "지적", it.get("reviewer"), it.get("stratum"), it.get("kind"),
                       it.get("as_is"), it.get("to_be"), it.get("reason"), it.get("severity"), it.get("confidence"),
                       it.get("선택"), it.get("반영")])
            row = ws[ws.max_row]
            fill = OK_FILL if it.get("kind") == "이상무" else ISSUE_FILL.get(it.get("severity"), None)
            if fill:
                for c in row:
                    c.fill = fill
        dv = DataValidation(type="list", formula1='"1안 적용,기각,보류"', allow_blank=True)
        ws.add_data_validation(dv)
        dv.add("N%d:N%d" % (hdr + 1, ws.max_row))
        ws.append([])
        ws.append(["집필·검토 패널 (가상 인물 — MatrAIx Persona-1M 표본)"]); ws.cell(ws.max_row, 1).font = Font(bold=True)
        ws.append(["번호", "구분", "자리", "가명", "약력"])
        for c in ws[ws.max_row]:
            c.font = HEAD_FONT; c.fill = HEAD_FILL
        for p_ in r.get("panel", []):
            ws.append([p_.get("번호"), p_.get("구분"), p_.get("자리"), p_.get("가명"), p_.get("약력")])
        widths = {"A": 6, "B": 16, "C": 8, "D": 26, "E": 8, "F": 16, "G": 11, "H": 8, "I": 45, "J": 45, "K": 40, "L": 7, "M": 7, "N": 10, "O": 24}
        for k, w in widths.items():
            ws.column_dimensions[k].width = w
        for row in ws.iter_rows(min_row=hdr + 1):
            for c in row:
                c.alignment = Alignment(vertical="top", wrap_text=True)
        ws.freeze_panes = "A%d" % (hdr + 1)
        made += 1
    wb.save(DATA / "검토의견.xlsx")
    return made


def roster_update(slots, stats):
    p = DATA / "00_패널" / "패널명부.xlsx"
    if not p.is_file():
        return
    wb = openpyxl.load_workbook(p)
    ws = wb["명부"]
    cols = [c.value for c in ws[1]]
    if "참여 강의" not in cols:
        return
    ci = cols.index("참여 강의") + 1
    ni = cols.index("번호") + 1
    part = {}
    for s in slots:
        lid = s["차시ID"]
        rv = LECT / lid / "08_검토.json"
        wr = LECT / lid / ("%s_원고.html" % lid)
        if rv.is_file():
            for p_ in json.load(io.open(rv, encoding="utf-8")).get("panel", []):
                part.setdefault(p_.get("번호"), []).append(lid + "(검토)" if p_.get("구분") == "검토진" else lid + "(집필)")
        elif wr.is_file():
            for row in ws.iter_rows(min_row=2, values_only=True):
                if row[cols.index("구분")] == "저자진":
                    part.setdefault(row[ni - 1], []).append(lid + "(집필)")
    for row in ws.iter_rows(min_row=2):
        no = row[ni - 1].value
        row[ci - 1].value = ", ".join(part.get(no, []))
    wb.save(p)


def toc_xlsx(slots, stats):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "목차"
    cols = ["배치", "과정", "모듈코드", "모듈명", "코드", "교과목명", "차시ID", "차수", "총차수", "강 제목", "학습목표 1", "학습목표 2", "학습목표 3", "단계", "장수", "낭독자", "분", "검수"]
    ws.append(cols)
    for c in ws[1]:
        c.font = HEAD_FONT; c.fill = HEAD_FILL
    for s in slots:
        st = stats[s["차시ID"]]
        ob = (st.get("학습목표") or ["", "", ""]) + ["", "", ""]
        stage = "탈고" if st["탈고"] else "검토" if st["검토"] else "초고" if st["초고"] else "스토리보드" if st["스토리보드"] else "설계" if st["설계"] else ""
        ws.append([s["배치"], s["과정"], s["모듈코드"], s["모듈명"], s["코드"], s["교과목명"], s["차시ID"], s["차수"], s["총차수"],
                   st["제목"] or s.get("제목") or "", ob[0], ob[1], ob[2], stage, st["장수"], st["글자수"], st["분"], st["판정"]])
    for i, k in enumerate(cols, 1):
        ws.column_dimensions[get_column_letter(i)].width = 40 if k.startswith("학습목표") or k in ("강 제목",) else 12
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(vertical="top", wrap_text=True)
    ws.freeze_panes = "H2"
    wb.save(DATA / "목차.xlsx")


def main():
    subjects, slots = load_input()
    stats = {s["차시ID"]: status_of(s["차시ID"]) for s in slots}
    io.open(DATA / "index.html", "w", encoding="utf-8").write(index_html(slots, stats))
    n = review_xlsx(slots, stats)
    roster_update(slots, stats)
    toc_xlsx(slots, stats)
    done = sum(1 for st in stats.values() if st["초고"])
    print("index.html · 검토의견.xlsx(%d강) · 목차.xlsx · 패널명부 갱신 — 초고 %d / 탈고 %d / %d강" % (
        n, done, sum(1 for st in stats.values() if st["탈고"]), len(slots)))


if __name__ == "__main__":
    main()
